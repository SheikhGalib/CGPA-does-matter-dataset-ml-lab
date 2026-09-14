"""Update presentation/What_Shapes_Student_CGPA_Ready.pptx with the final results (920 students,
5-fold CV, result satisfaction not used) and add the data-check slides.

Always starts from the untouched backup, so it can be re-run safely:
    presentation/What_Shapes_Student_CGPA_Ready_BEFORE_FINAL_UPDATE.pptx
Needs: docs/final_results.json (notebooks/final_weka_results.py) and docs/figures/final/ (make_figures.py).
"""
import copy
import json
import re
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parents[2]
DECK = ROOT / "presentation" / "What_Shapes_Student_CGPA_Ready.pptx"
BACKUP = ROOT / "presentation" / "What_Shapes_Student_CGPA_Ready_BEFORE_FINAL_UPDATE.pptx"
FIG = ROOT / "docs" / "figures" / "final"
R = json.loads((ROOT / "docs" / "final_results.json").read_text(encoding="utf-8"))
if not BACKUP.exists():
    shutil.copy2(DECK, BACKUP)
prs = Presentation(BACKUP)
S = {i + 1: s for i, s in enumerate(prs.slides)}          # original slide numbers
N = R["final_rows"]
CV = {t: R["weka"][t]["cv5"] for t in R["weka"]}
SPLIT = {t: R["weka"][t]["split80"] for t in R["weka"]}
BAND_RGB = {"C0": "A83A3A", "C1": "DC713D", "C2": "4C91D8", "C3": "285B9F"}


# ============================================================================ helpers
def texty(sh):
    return getattr(sh, "has_text_frame", False) and sh.has_text_frame


def find(slide, start=None, exact=None, contains=None, name=None):
    for sh in slide.shapes:
        if name is not None and sh.name != name:
            continue
        if name is not None and start is None and exact is None and contains is None:
            return sh
        if not texty(sh):
            continue
        t = sh.text_frame.text.strip()
        if exact is not None and t == exact:
            return sh
        if start is not None and t.startswith(start):
            return sh
        if contains is not None and contains in t:
            return sh
    raise KeyError(f"shape not found: start={start!r} exact={exact!r} contains={contains!r} name={name!r}")


def _set_para(p, text):
    for br in p._p.findall(qn("a:br")):
        p._p.remove(br)
    runs = p.runs
    if not runs:
        p.add_run().text = text
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def set_text(shape_or_tf, text):
    """Replace text, keeping the formatting of each existing paragraph's first run."""
    tf = shape_or_tf.text_frame if hasattr(shape_or_tf, "text_frame") else shape_or_tf
    lines = text.split("\n")
    paras = list(tf.paragraphs)
    while len(paras) < len(lines):
        paras[-1]._p.addnext(copy.deepcopy(paras[-1]._p))
        paras = list(tf.paragraphs)
    for p, line in zip(paras, lines):
        _set_para(p, line)
    for p in paras[len(lines):]:
        p._p.getparent().remove(p._p)


def set_color(shape, hexrgb):
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor.from_string(hexrgb)


def replace_picture(slide, pic, png):
    _, rid = slide.part.get_or_add_image_part(str(png))
    blip_fill = pic._element.blipFill
    blip_fill.blip.set(qn("r:embed"), rid)
    src = blip_fill.find(qn("a:srcRect"))
    if src is not None:
        blip_fill.remove(src)


def pictures(slide):
    return sorted([s for s in slide.shapes if s.shape_type == 13], key=lambda s: (s.left, s.top))


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def notes_replace(slide, pairs):
    if not slide.has_notes_slide:
        return
    tf = slide.notes_slide.notes_text_frame
    t = tf.text
    for a, b in pairs:
        t = t.replace(a, b)
    tf.text = t


def table_of(slide):
    return next(s for s in slide.shapes if getattr(s, "has_table", False) and s.has_table).table


def set_cell(table, r, c, text):
    set_text(table.cell(r, c).text_frame, text)


def pct(x, d=2):
    return f"{100 * x:.{d}f}%"


def update_cms(slide, matrices):
    cells = [s for s in slide.shapes if texty(s) and re.fullmatch(r"\d+", s.text_frame.text.strip())
             and s.width < Pt(70) and Pt(185) < s.top < Pt(470)]
    cells.sort(key=lambda s: s.left)
    groups, cur = [], [cells[0]]
    for s in cells[1:]:
        if s.left - cur[-1].left > Pt(120):
            groups.append(cur)
            cur = [s]
        else:
            cur.append(s)
    groups.append(cur)
    assert len(groups) == len(matrices) and all(len(g) == 16 for g in groups), [len(g) for g in groups]
    for g, m in zip(groups, matrices):
        g.sort(key=lambda s: (round(s.top / 12700), s.left))
        for k, s in enumerate(g):
            set_text(s, str(m[k // 4][k % 4]))


def fill_hex(sh):
    try:
        return str(sh.fill.fore_color.rgb) if sh.fill.type == 1 else None
    except Exception:
        return None


def update_bar_slide(slide, item):
    mix = {r["code"]: r for r in R["per_question_cgpa_mix"][item]}
    bars = [s for s in slide.shapes if fill_hex(s) in set(BAND_RGB.values()) and s.height > Pt(25)]
    labels = [s for s in slide.shapes if texty(s) and re.fullmatch(r"\d+%", s.text_frame.text.strip())]
    ns = [s for s in slide.shapes if texty(s) and s.text_frame.text.strip().startswith("n=")]
    tops = sorted({b.top for b in bars})
    rows = []
    for t in tops:
        if rows and abs(t - rows[-1]) < Pt(3):
            continue
        rows.append(t)
    assert len(rows) == 5, (item, len(rows))
    for code, top in enumerate(rows):
        rb = sorted([b for b in bars if abs(b.top - top) < Pt(3)], key=lambda b: b.left)
        rl = sorted([l for l in labels if top - Pt(2) <= l.top <= top + Pt(30)], key=lambda l: l.left)
        rn = [n for n in ns if top - Pt(2) <= n.top <= top + Pt(30)]
        assert len(rb) == 4 and len(rl) == 4 and len(rn) == 1, (item, code, len(rb), len(rl), len(rn))
        x0, total = rb[0].left, rb[-1].left + rb[-1].width - rb[0].left
        row = mix.get(code)
        shares = [row[f"C{k}"] for k in range(4)] if row else [0, 0, 0, 0]
        ssum = sum(shares) or 1
        x = x0
        for k, (b, l) in enumerate(zip(rb, rl)):
            w = int(round(total * shares[k] / ssum)) if k < 3 else x0 + total - x
            w = max(w, 0) if row else 0
            b.left, b.width = int(x), int(w)
            l.left, l.width = int(x + Pt(3)), int(max(w - Pt(4.5), Pt(1)))
            set_text(l, f"{shares[k]}%" if row and shares[k] >= 6 else "")
            x += w
        set_text(rn[0], f"n={row['n'] if row else 0}")


def duplicate_slide(src):
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    for el in src.shapes._spTree.iterchildren():
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        new.shapes._spTree.append(copy.deepcopy(el))
    src_bg = src._element.cSld.find(qn("p:bg"))
    if src_bg is not None:
        new._element.cSld.insert(0, copy.deepcopy(src_bg))
    for el in new.shapes._spTree.iter():
        for attr in (qn("r:embed"), qn("r:link"), qn("r:id")):
            rid = el.get(attr)
            if rid and rid in src.part.rels:
                rel = src.part.rels[rid]
                if rel.is_external:
                    new_rid = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
                else:
                    new_rid = new.part.relate_to(rel.target_part, rel.reltype)
                el.set(attr, new_rid)
    return new


def move_to_after(slide, anchor):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    sid = next(i for i in ids if prs.part.related_part(i.rId) is slide.part)
    aid = next(i for i in ids if prs.part.related_part(i.rId) is anchor.part)
    lst.remove(sid)
    aid.addnext(sid)


def hide(slide):
    slide._element.set("show", "0")


# ============================================================================ numbers
ans = {k: {r["value"]: r["n"] for r in v} for k, v in R["answers"].items()}
def share(item, values):
    n = sum(ans[item].get(v, 0) for v in values)
    return n, 100 * n / N

G = R["gap_analysis"]
assoc_c, assoc_s = R["association_cgpa"], R["association_sgpa"]

# ============================================================================ duplicates first (from pristine slides)
NEW_AGREE = duplicate_slide(S[15])
NEW_EXAMPLE = duplicate_slide(S[65])
NEW_WHERE = duplicate_slide(S[18])
NEW_WHEN = duplicate_slide(S[13])
NEW_RANDOM = duplicate_slide(S[18])
NEW_KEPT = duplicate_slide(S[9])
NEW_BEFORE_AFTER = duplicate_slide(S[36])

# ============================================================================ existing slides
# --- 1 cover
set_text(find(S[1], start="1,108 usable"), "920 checked responses from four universities\nBehaviour patterns, SGPA prediction, and CGPA prediction")
notes_replace(S[1], [("1,108 jon", "920 jon")])

# --- 3 data collection
set_text(find(S[3], exact="1,108"), "920")
set_text(find(S[3], exact="usable records"), "final records")
notes_replace(S[3], [("Clean korar por 1,108 ta usable record.", "Clean ar SGPA-CGPA check er por 920 ta final record.")])

# --- 8 (hidden) replace the emoji icons with plain numbers
ovals = sorted([s for s in S[8].shapes if s.name.startswith("Oval")], key=lambda s: (s.top, s.left))
for k, sh in enumerate(ovals):
    set_text(sh, str(k + 1))

# --- 13 universities
replace_picture(S[13], pictures(S[13])[0], FIG / "rc_university.png")
set_text(find(S[13], exact="1,108"), f"{N}")
set_text(find(S[13], exact="usable students"), "students in the final dataset")
u = {r["value"]: r for r in R["universities"]}
set_notes(S[13], f"After the data check we have {N} students. The universities are not equal: KUET {u['KUET']['n']} "
                 f"({u['KUET']['pct']}%), BUET {u['BUET']['n']}, BRAC {u['BRAC']['n']}, CUET {u['CUET']['n']}. "
                 "So we cannot say the results apply equally to every university.")

# --- 14 four CGPA groups
bands = {b["band"]: b for b in R["cgpa_bands"]}
for old_n, old_p, band in [("236", "students · 21.3%", "C0"), ("321", "students · 29.0%", "C1"),
                           ("303", "students · 27.3%", "C2"), ("248", "students · 22.4%", "C3")]:
    set_text(find(S[14], exact=old_n), str(bands[band]["n"]))
    set_text(find(S[14], exact=old_p), f"students · {bands[band]['pct']:.1f}%")
set_notes(S[14], "The form collected CGPA ranges, not exact numbers, so we made four groups: "
                 f"C0 below 3.20 ({bands['C0']['n']} students), C1 3.20 to 3.49 ({bands['C1']['n']}), "
                 f"C2 3.50 to 3.74 ({bands['C2']['n']}), C3 3.75 or above ({bands['C3']['n']}). "
                 "This is a four-class classification problem. These C0-C3 labels stay the same on every slide.")

# --- 15 CGPA distribution
replace_picture(S[15], pictures(S[15])[0], FIG / "rc_cgpa.png")
largest = max(R["cgpa_bands"], key=lambda b: b["n"])
smallest = min(R["cgpa_bands"], key=lambda b: b["n"])
set_text(find(S[15], exact="29.0%"), f"{largest['pct']:.1f}%")
set_text(find(S[15], exact="largest (C1) = guessing baseline"), f"largest ({largest['band']}) = guessing baseline")
set_text(find(S[15], exact="21.3%"), f"{smallest['pct']:.1f}%")
set_text(find(S[15], exact="smallest group (C0)"), f"smallest group ({smallest['band']})")
set_notes(S[15], f"Before building models we checked that no group dominates. The largest group is {largest['band']} "
                 f"({largest['pct']:.1f}%), the smallest is {smallest['band']} ({smallest['pct']:.1f}%). "
                 "The four groups are roughly balanced, so accuracy is a fair measure.")

# --- 16 attendance
replace_picture(S[16], pictures(S[16])[0], FIG / "rc_attendance.png")
most = max(R["answers"]["attendance"], key=lambda r: r["n"])
n75, p75 = share("attendance", ["75-89%", "90% or more"])
n60, p60 = share("attendance", ["Less than 40%", "40-59%"])
set_text(find(S[16], exact="418 · 37.7%"), f"{most['n']} · {most['pct']:.1f}%")
set_text(find(S[16], exact="most common: 90% or more"), f"most common: {most['value']}")
set_text(find(S[16], exact="799 · 72.1%"), f"{n75} · {p75:.1f}%")
set_text(find(S[16], exact="33 · 3.0%"), f"{n60} · {p60:.1f}%")
set_notes(S[16], f"This slide only shows answers, no prediction. The most common answer is '{most['value']}' ({most['n']} students). "
                 f"{p75:.1f}% attend 75% or more, and only {n60} attend below 60%.")

# --- 17 study time
replace_picture(S[17], pictures(S[17])[0], FIG / "rc_weekly_study_time.png")
most = max(R["answers"]["weekly_study_time"], key=lambda r: r["n"])
nlt3, plt3 = share("weekly_study_time", ["Less than 3 hrs"])
ngt10, pgt10 = share("weekly_study_time", ["11-15 hrs", "More than 15 hrs"])
set_text(find(S[17], exact="281 · 25.4%"), f"{most['n']} · {most['pct']:.1f}%")
set_text(find(S[17], exact="most common: 3–6 hrs a week"), f"most common: {most['value'].replace('-', '–')} a week")
set_text(find(S[17], exact="235 · 21.2%"), f"{nlt3} · {plt3:.1f}%")
set_text(find(S[17], exact="341 · 30.8%"), f"{ngt10} · {pgt10:.1f}%")
set_notes(S[17], f"Study time outside class. Most students study {most['value']} a week ({most['n']} students). "
                 f"{nlt3} study less than 3 hours, and {ngt10} study more than 10 hours.")

# --- 18 stress + sleep
left, right = pictures(S[18])
replace_picture(S[18], left, FIG / "rc_stress_frequency.png")
replace_picture(S[18], right, FIG / "rc_sleep_duration.png")
ns_, ps_ = share("stress_frequency", ["Often", "Almost always"])
nsl, psl = share("sleep_duration", ["Less than 5 hrs", "5-6 hrs"])
set_text(find(S[18], start="Often + almost always"), f"Often + almost always: {ns_} students ({ps_:.1f}%)")
set_text(find(S[18], start="Under 6 hours"), f"Under 6 hours: {nsl} students ({psl:.1f}%)")
set_notes(S[18], f"{ps_:.1f}% of students say stress often or almost always made studying difficult. "
                 f"{psl:.1f}% sleep less than 6 hours. Here we only describe answers; links with results come later.")

# --- 19 distraction + commitments
left, right = pictures(S[19])
replace_picture(S[19], left, FIG / "rc_distraction_frequency.png")
replace_picture(S[19], right, FIG / "rc_weekly_responsibilities.png")
nd, pd_ = share("distraction_frequency", ["Often", "Almost every day"])
nr, pr = share("weekly_responsibilities", ["6-10 hrs", "11-15 hrs", "More than 15 hrs"])
set_text(find(S[19], start="Often + almost every day"), f"Often + almost every day: {nd} ({pd_:.1f}%)")
set_text(find(S[19], start="6 hours or more"), f"6 hours or more: {nr} ({pr:.1f}%)")
set_notes(S[19], f"{pd_:.1f}% say phone or social media often or almost every day took their study time. "
                 f"{pr:.1f}% spend 6 or more hours a week on tuition, job or club work.")

# --- 20 motivation, support, environment
pics = pictures(S[20])
for pic, item in zip(pics, ["desired_department_match", "support_level", "study_environment"]):
    replace_picture(S[20], pic, FIG / f"rc_{item}_small.png")
n1, p1 = share("desired_department_match", ["Yes, exactly what I wanted"])
n2, p2 = share("support_level", ["Good support", "Strong support"])
n3, p3 = share("study_environment", ["Mostly suitable", "Excellent for focused study"])
set_text(find(S[20], start="Exactly what I wanted"), f"Exactly what I wanted: {n1} ({p1:.1f}%)")
set_text(find(S[20], start="Good + strong"), f"Good + strong: {n2} ({p2:.1f}%)")
set_text(find(S[20], start="Mostly suitable + excellent"), f"Mostly suitable + excellent: {n3} ({p3:.1f}%)")
set_notes(S[20], f"{p1:.1f}% got exactly the department they wanted. {p2:.1f}% get good or strong support from family "
                 f"or friends, and {p3:.1f}% say their living place is mostly suitable or excellent for study.")

# --- 21 associations
replace_picture(S[21], pictures(S[21])[0], FIG / "association_strength.png")
set_text(find(S[21], start="small positive pattern (+0.12)"), f"small positive pattern ({assoc_c['desired_department_match']['rho']:+.2f})")
set_text(find(S[21], start="small positive pattern (+0.09)"), f"small positive pattern ({assoc_c['attendance']['rho']:+.2f})")
set_text(find(S[21], start="small negative pattern"), f"small negative pattern ({assoc_c['weekly_responsibilities']['rho']:+.2f})")
reliable = [k for k, v in assoc_c.items() if v["reliable"]]
set_notes(S[21], f"Out of 12 behaviour questions, {len(reliable)} patterns were reliable after checking: getting the desired "
                 "department and higher attendance go with a higher CGPA group, and more tuition/job/club hours go with a "
                 f"lower group. All links are small (the largest is {max(abs(v['rho']) for v in assoc_c.values()):.2f}), "
                 "so we say 'associated', not 'causes'.")

# --- 22 SGPA vs CGPA patterns
replace_picture(S[22], pictures(S[22])[0], FIG / "sgpa_vs_cgpa_pattern.png")
set_text(find(S[22], start="✔ Class attendance"), f"Class attendance\nsame direction: SGPA {assoc_s['attendance']['rho']:+.2f}, CGPA {assoc_c['attendance']['rho']:+.2f}")
set_text(find(S[22], start="✔ Got desired department"), f"Got desired department\nsame direction: SGPA {assoc_s['desired_department_match']['rho']:+.2f}, CGPA {assoc_c['desired_department_match']['rho']:+.2f}")
set_notes(S[22], "We checked whether each answer moves in the same direction for SGPA and CGPA. Attendance and getting the "
                 "desired department go the same way for both. The other 10 questions are too weak to claim a matching pattern.")

# --- 23-34 per-question CGPA bars
BAR_SLIDES = {23: "attendance", 24: "weekly_study_time", 25: "study_style", 26: "topic_clarity", 27: "sleep_duration",
              28: "stress_frequency", 29: "distraction_frequency", 30: "weekly_responsibilities",
              31: "desired_department_match", 32: "support_level", 33: "study_environment", 34: "routine_manageability"}
for sn, item in BAR_SLIDES.items():
    update_bar_slide(S[sn], item)
    set_notes(S[sn], f"Source: docs/final_results.json ({N} students after the data check). Each bar = 100% of the students who chose that answer.")
set_text(find(S[29], exact="Attendance"), "Phone Distraction")
set_text(find(S[29], contains="N = 1,108"), f"Each bar = 100% of students choosing that answer. Descriptive association, N = {N:,}.")

# --- 35 (hidden) framework: remove emojis, new split sizes
ovals = sorted([s for s in S[35].shapes if s.name.startswith("Oval")], key=lambda s: s.left)
for k, sh in enumerate(ovals):
    set_text(sh, str(k + 1))
n_train = N - SPLIT["main_cgpa"]["J48"]["n"]
set_text(find(S[35], exact="881 (80%)"), f"{n_train} (80%)")
set_text(find(S[35], exact="221 (20%)"), f"{SPLIT['main_cgpa']['J48']['n']} (20%)")

# --- 36 WEKA results table (5-fold CV)
tb = table_of(S[36])
for r, m in enumerate(["J48", "RandomForest", "SMO", "Logistic", "NaiveBayes", "ZeroR"], start=1):
    for c, t in enumerate(["sgpa", "questionnaire_cgpa", "main_cgpa"], start=1):
        set_cell(tb, r, c, pct(CV[t][m]["accuracy"]))
set_text(find(S[36], start="Stratified 10-fold"), "Stratified 5-fold cross-validation, seed 1")
set_text(find(S[36], start="Same 1,102"), f"Same {N} students. Questionnaire: 12 inputs. Main CGPA: same inputs + recent SGPA.")
set_notes(S[36], "Source: WEKA 3.8.7 logs in docs/weka_runs/final/ (files ending _cv5.txt). Result satisfaction is not used as an input.")

# --- 37 / 39 / 41 trees
for sn, t, label in [(37, "sgpa", "Recent SGPA"), (39, "questionnaire_cgpa", "CGPA from questionnaire variables"),
                     (41, "main_cgpa", "Main CGPA prediction")]:
    replace_picture(S[sn], pictures(S[sn])[0], FIG / f"tree_{t}.png")
    set_text(find(S[sn], start="WEKA J48 -C"),
             f"WEKA J48 -C 0.25 -M 40, built on all {N} students. {label}. 5-fold CV: {pct(CV[t]['J48']['accuracy'])}. "
             f"ZeroR: {pct(CV[t]['ZeroR']['accuracy'])}.")
    set_notes(S[sn], f"Tree drawn from the WEKA J48 output (docs/weka_runs/final/{t}_J48_tree_full.txt). "
                     "Leaf (N/errors) = students in that leaf / students it places in the wrong group.")

# --- 38 (hidden) SGPA rules
c1, c2 = find(S[38], start="IF routine is completely"), find(S[38], start="IF routine is above")
set_text(c1, "IF the student did not get the desired department\nAND phone does not take study time almost every day")
set_text(c2, "IF got the desired department AND attendance 90% or more\nAND self-study is 7 or more hours a week")
k1, k2 = find(S[38], exact="C1"), find(S[38], exact="C0")
set_text(k1, "C1"); set_color(k1, BAND_RGB["C1"])
set_text(k2, "C3"); set_color(k2, BAND_RGB["C3"])
set_notes(S[38], "Rules come from the WEKA J48 tree in docs/weka_runs/final/sgpa_J48_tree_full.txt. They are predictions, not causes.")

# --- 40 questionnaire rules
c1, c2 = find(S[40], start="IF study environment is very"), find(S[40], start="IF environment is excellent")
set_text(c1, "IF the living place is very distracting\nAND the student did not get the desired department")
set_text(c2, "IF the living place is not very distracting\nAND stress almost always made studying difficult")
k1, k2 = find(S[40], exact="C0"), find(S[40], exact="C3")
set_text(k2, "C2"); set_color(k2, BAND_RGB["C2"])
set_notes(S[40], "Rules come from the WEKA J48 tree in docs/weka_runs/final/questionnaire_cgpa_J48_tree_full.txt. They are predictions, not causes.")

# --- 42 main CGPA rules (the tree is now exactly the four SGPA rules)
set_text(find(S[42], start="IF recent SGPA is 3.75 or higher"), "IF recent SGPA is 3.75 or higher")
set_text(find(S[42], start="Selected branches"), "The whole WEKA J48 tree: four rules, one for each recent SGPA group.")
set_notes(S[42], "Source: docs/weka_runs/final/main_cgpa_J48_tree_full.txt. With at least 40 students per leaf, J48 keeps only recent SGPA.")

# --- 43 (hidden) J48 CV and test accuracy
tb = table_of(S[43])
set_cell(tb, 0, 1, f"5-fold CV ({N})")
set_cell(tb, 0, 2, f"80/20 test ({SPLIT['main_cgpa']['J48']['n']})")
set_cell(tb, 0, 3, "ZeroR test")
for r, t in enumerate(["sgpa", "questionnaire_cgpa", "main_cgpa"], start=1):
    set_cell(tb, r, 1, pct(CV[t]["J48"]["accuracy"]))
    set_cell(tb, r, 2, pct(SPLIT[t]["J48"]["accuracy"]))
    set_cell(tb, r, 3, pct(SPLIT[t]["ZeroR"]["accuracy"]))
set_text(find(S[43], name="Google Shape;533;p58"), "WEKA J48: cross-validation and 80/20 test accuracy")
set_text(find(S[43], start="Each row evaluates"), f"J48 with minNumObj 40. 5-fold CV on all {N} students; 80/20 = WEKA percentage split, seed 1 ({n_train} train, {SPLIT['main_cgpa']['J48']['n']} test).")

# --- 44 (hidden) J48 80/20 confusion matrix
update_cms(S[44], [SPLIT["main_cgpa"]["J48"]["confusion_matrix"]])
set_text(find(S[44], exact="Fixed test set"), f"80/20 split, {SPLIT['main_cgpa']['J48']['n']} test students")
set_notes(S[44], "Source: docs/weka_runs/final/main_cgpa_J48_split80.txt (WEKA percentage split 80%, seed 1).")

# --- 45 / 46 / 47 confusion matrices (5-fold CV)
update_cms(S[45], [CV["main_cgpa"]["RandomForest"]["confusion_matrix"], CV["main_cgpa"]["SMO"]["confusion_matrix"]])
update_cms(S[46], [CV["main_cgpa"]["Logistic"]["confusion_matrix"], CV["main_cgpa"]["NaiveBayes"]["confusion_matrix"]])
for sn in (45, 46):
    set_text(find(S[sn], start="Main CGPA, stratified"), f"Main CGPA, stratified 5-fold cross-validation, N = {N}. Rows actual, columns predicted.")
set_notes(S[45], "Sources: docs/weka_runs/final/main_cgpa_RandomForest_cv5.txt and main_cgpa_SMO_cv5.txt.")
set_notes(S[46], "Sources: docs/weka_runs/final/main_cgpa_Logistic_cv5.txt and main_cgpa_NaiveBayes_cv5.txt.")
update_cms(S[47], [CV["main_cgpa"]["ZeroR"]["confusion_matrix"]])
set_text(find(S[47], exact="Main CGPA, 10-fold CV"), "Main CGPA, 5-fold CV")
set_text(find(S[47], start="WEKA ZeroR, main CGPA"), f"WEKA ZeroR, main CGPA, 5-fold CV, N = {N}. Accuracy = {pct(CV['main_cgpa']['ZeroR']['accuracy'])}.")
set_notes(S[47], "Source: docs/weka_runs/final/main_cgpa_ZeroR_cv5.txt.")

# --- 48 accuracy chart
chart_shape = next(s for s in S[48].shapes if s.has_chart)
chart = chart_shape.chart
cats = ["ZeroR", "J48", "Random Forest", "SMO", "Logistic", "Naive Bayes"]
keys = ["ZeroR", "J48", "RandomForest", "SMO", "Logistic", "NaiveBayes"]
cd = CategoryChartData()
cd.categories = cats
cd.add_series(chart.plots[0].series[0].name or "Accuracy",
              [round(100 * CV["main_cgpa"][k]["accuracy"], 2) for k in keys], number_format='0"%"')
chart.replace_data(cd)
va = chart.value_axis
va.maximum_scale, va.minimum_scale = 60, 0
va.tick_labels.number_format, va.tick_labels.number_format_is_linked = '0"%"', False
dl = chart.plots[0].data_labels
dl.number_format, dl.number_format_is_linked = '0.0"%"', False
for dlbls in chart._chartSpace.xpath('.//c:dLbls'):
    nf = dlbls.find(qn('c:numFmt'))
    if nf is None:
        nf = dlbls.makeelement(qn('c:numFmt'), {})
        kids = [k for k in dlbls if k.tag != qn('c:dLbl')]
        (kids[0].addprevious(nf) if kids else dlbls.append(nf))
    nf.set('formatCode', '0.0"%"')
    nf.set('sourceLinked', '0')
set_text(find(S[48], name="Google Shape;533;p58"), "WEKA main CGPA accuracy: 5-fold CV")
set_text(find(S[48], start="WEKA 3.8.7, stratified"), f"WEKA 3.8.7, stratified 5-fold CV, seed 1. Same {N} students and main CGPA inputs.")
set_notes(S[48], "Sources: docs/weka_runs/final/main_cgpa_*_cv5.txt")

# --- 50 / 51 / 52 / 53
set_text(find(S[50], start="Engineering students, four CGPA bands"),
         "Engineering students, four CGPA bands\nJ48, Random Forest, SMO, Logistic,\nNaive Bayes and ZeroR\nWEKA 5-fold comparison")
before = R["before_check_1102_cv5"]["main"]
set_text(find(S[51], start="Student answers vary"),
         f"Checking SGPA against CGPA removed {G['dropped_rows']} unreliable answers.")
set_text(find(S[51], start="Random Forest leads"),
         f"Accuracy rose from about {100 * before['SMO']['accuracy']:.0f}% to {100 * CV['main_cgpa']['SMO']['accuracy']:.0f}% (SMO); J48 gives readable rules.")
set_notes(S[51], "Sources: docs/final_results.json and docs/weka_runs/final/.")
set_text(find(S[52], start="KUET-heavy"), f"KUET-heavy sample ({u['KUET']['pct']:.0f}%) and a one-time survey")
set_text(find(S[52], start="Questionnaire imputation"), f"Data check removed {G['dropped_rows']} answers; some kept answers may still be careless")
set_text(find(S[52], start="Small associations"), "Small associations; prediction relies mostly on recent SGPA")
set_notes(S[52], "The data check only looked at the 6-7 August KUET semester-7 group. Other students with a big SGPA/CGPA gap (for example 37 at BUET) were kept.")
set_text(find(S[53], start="1,108 usable"), f"{N} checked engineering-student responses (from 1,126 collected)")

# --- 54 WEKA output
set_text(find(S[54], start="We Also Ran the Models"), "We Also Ran the Models in WEKA")
left, right = pictures(S[54])
replace_picture(S[54], left, FIG / "weka_output_main_rf_cv5.png")
replace_picture(S[54], right, FIG / "weka_output_main_smo_cv5.png")
set_text(find(S[54], start="Main CGPA Random Forest output"), f"Main CGPA Random Forest output — {N} instances, 5-fold cross-validation")
set_text(find(S[54], start="SMO on the main CGPA task"),
         f"SMO on the main CGPA task, 5-fold cross-validation: {100 * CV['main_cgpa']['SMO']['accuracy']:.1f}% correct · confusion matrix rows = actual")
set_text(find(S[54], exact="46.5%"), f"{100 * CV['main_cgpa']['SMO']['accuracy']:.1f}%")
set_text(find(S[54], exact="46.3%"), f"{100 * CV['main_cgpa']['RandomForest']['accuracy']:.1f}%")
set_text(find(S[54], exact="29.0%"), f"{100 * CV['main_cgpa']['ZeroR']['accuracy']:.1f}%")
set_notes(S[54], "These panels show the real WEKA 3.8.7 output text. Replace them with your own WEKA Explorer screenshots "
                 "(steps in docs/FINAL_WEKA_STEPS.md). Main CGPA SMO: "
                 f"{100 * CV['main_cgpa']['SMO']['accuracy']:.1f}%, Random Forest {100 * CV['main_cgpa']['RandomForest']['accuracy']:.1f}%, "
                 f"ZeroR (always guessing the most common group) {100 * CV['main_cgpa']['ZeroR']['accuracy']:.1f}%.")

# --- 55 (hidden) previous work table, last row
tb = table_of(S[55])
last = len(tb.rows) - 1
set_cell(tb, last, 1, f"{N} engineering students (after the data check), 4 universities: CGPA group C0–C3")
set_cell(tb, last, 2, "WEKA J48 for rules · Random Forest, SMO, Logistic, Naive Bayes, ZeroR")
set_cell(tb, last, 3, f"WEKA SMO {100 * CV['main_cgpa']['SMO']['accuracy']:.1f}% (5-fold CV) · J48 {100 * SPLIT['main_cgpa']['J48']['accuracy']:.1f}% on an 80/20 test split")

# --- 60 (hidden) data dictionary
tb = table_of(S[60])
for r in range(len(tb.rows)):
    key = tb.cell(r, 0).text.strip()
    if key == "recent_sgpa":
        set_cell(tb, r, 6, "6 rows without a real SGPA removed")
    elif key == "cgpa":
        set_cell(tb, r, 6, f"18 no-CGPA + {G['dropped_rows']} big-gap rows removed")
    elif key == "result_satisfaction":
        set_cell(tb, r, 6, "Not used: depends on personal goals (subjective)")

# --- 61 (hidden) split counts and tree settings
tb1, tb2 = [s.table for s in S[61].shapes if getattr(s, "has_table", False) and s.has_table]
def counts(t, e):
    return [sum(row) for row in R["weka"][t][e]["J48"]["confusion_matrix"]]
for r, (label, t) in enumerate([("CGPA", "main_cgpa"), ("SGPA", "sgpa")]):
    test = counts(t, "split80")
    full = [b["n"] for b in (R["cgpa_bands"] if t == "main_cgpa" else R["sgpa_bands"])]
    train = [f - s for f, s in zip(full, test)]
    for k, v in enumerate(train):
        set_cell(tb1, 1 + 2 * r, 1 + k, str(v))
    set_cell(tb1, 1 + 2 * r, 5, str(sum(train)))
    for k, v in enumerate(test):
        set_cell(tb1, 2 + 2 * r, 1 + k, str(v))
    set_cell(tb1, 2 + 2 * r, 5, str(sum(test)))
set_cell(tb2, 0, 4, "5-fold CV accuracy")
sizes = R["j48_sizes"]
for r, t in enumerate(["sgpa", "questionnaire_cgpa", "main_cgpa"], start=1):
    set_cell(tb2, r, 1, str(sizes[f"{t}_M40"]["depth"]))
    set_cell(tb2, r, 2, "40")
    set_cell(tb2, r, 3, str(sizes[f"{t}_M40"]["leaves"]))
    set_cell(tb2, r, 4, f"{100 * CV[t]['J48']['accuracy']:.1f}%")
    set_cell(tb2, r, 5, f"{100 * SPLIT[t]['J48']['accuracy']:.1f}%")
set_text(find(S[61], start="Tree size chosen"),
         f"WEKA J48 with at least 40 students per leaf (minNumObj 40). 5-fold CV on all {N} students; test = WEKA 80/20 percentage split, seed 1.")

# --- 62 (hidden) tree size
set_text(find(S[62], exact="315"), str(sizes["questionnaire_cgpa_M2"]["leaves"]))
set_text(find(S[62], exact="leaves: WEKA J48, CGPA diagnostic"), "leaves: default WEKA J48, CGPA from questions")
set_text(find(S[62], exact="268"), str(sizes["main_cgpa_M2"]["leaves"]))
set_text(find(S[62], exact="leaves: WEKA J48, main CGPA task"), "leaves: default WEKA J48, main CGPA")
lv = sorted(sizes[f"{t}_M40"]["leaves"] for t in ["sgpa", "questionnaire_cgpa", "main_cgpa"])
set_text(find(S[62], exact="6–8"), f"{lv[0]}–{lv[-1]}")
set_text(find(S[62], exact="final boxes in our shallow trees"), "leaves with at least 40 students per leaf")
set_text(find(S[62], start="Real WEKA J48 tree viewer"), "Default J48 tree (main CGPA): too big to read")
p16, p18, p20 = pictures(S[62])
replace_picture(S[62], p16, FIG / "tree_main_cgpa_default_M2.png")
replace_picture(S[62], p18, FIG / "cm_split80_J48_sgpa.png")
replace_picture(S[62], p20, FIG / "cm_split80_J48_questionnaire_cgpa.png")
set_notes(S[62], f"Default WEKA J48 grows {sizes['main_cgpa_M2']['leaves']} leaves for main CGPA and "
                 f"{sizes['questionnaire_cgpa_M2']['leaves']} for the questionnaire tree, which cannot be explained on a slide. "
                 "With at least 40 students per leaf the trees stay small. The two matrices are J48 on the 80/20 test split.")

# --- 63 (hidden) full WEKA results, 5-fold CV
tb = table_of(S[63])
set_text(find(S[63], start="Full WEKA Results"), "Full WEKA Results (5-fold cross-validation)")
for r in range(1, len(tb.rows)):
    m = tb.cell(r, 0).text.strip()
    for c, t in [(1, "main_cgpa"), (4, "questionnaire_cgpa")]:
        v = CV[t][m]
        set_cell(tb, r, c, f"{100 * v['accuracy']:.1f}%")
        set_cell(tb, r, c + 1, "—" if m == "ZeroR" else f"{v['weighted_f1']:.3f}")
        set_cell(tb, r, c + 2, f"{v['kappa']:.3f}")
set_text(find(S[63], start="Source: raw WEKA"), "Source: WEKA logs in docs/weka_runs/final/. W-F1 = weighted F1. Kappa = agreement beyond chance.")

# --- 64 (hidden) now: WEKA 80/20 test results
tb = table_of(S[64])
set_text(find(S[64], start="Full Python Results"), "WEKA Results on the 80/20 Test Split")
set_cell(tb, 0, 2, "W-F1")
set_cell(tb, 0, 5, "W-F1")
models_split = ["ZeroR", "OneR", "J48", "NaiveBayes", "Logistic", "SMO", "RandomForest"]
while len(tb.rows) - 1 > len(models_split):
    tr = tb._tbl.tr_lst[-1]
    tr.getparent().remove(tr)
for r, m in enumerate(models_split, start=1):
    set_cell(tb, r, 0, m)
    for c, t in [(1, "main_cgpa"), (4, "questionnaire_cgpa")]:
        v = SPLIT[t][m]
        set_cell(tb, r, c, f"{100 * v['accuracy']:.1f}%")
        set_cell(tb, r, c + 1, "—" if m == "ZeroR" else f"{v['weighted_f1']:.3f}")
        set_cell(tb, r, c + 2, f"{v['kappa']:.3f}")
set_notes(S[64], f"WEKA percentage split 80% (seed 1): {n_train} students for training, {SPLIT['main_cgpa']['J48']['n']} for testing. "
                 "Logs: docs/weka_runs/final/*_split80.txt. Show this only if the teacher asks; the main result is 5-fold CV.")

# --- 65 (hidden) CV vs hold-out
set_text(find(S[65], start="• One fixed 80/20"),
         f"• One fixed 80/20 split\n• {n_train} train · {SPLIT['main_cgpa']['J48']['n']} test students\n"
         f"• Main CGPA J48 tree: {100 * SPLIT['main_cgpa']['J48']['accuracy']:.1f}% on test\n• Easy to explain: rules → unseen students")
set_text(find(S[65], start="• Every student is tested"),
         f"• Every student is tested once\n• 5 folds, seed 1\n• Main CGPA SMO: {100 * CV['main_cgpa']['SMO']['accuracy']:.1f}% (J48 {100 * CV['main_cgpa']['J48']['accuracy']:.1f}%)\n"
         "• Less luck from one split")
set_notes(S[65], "Hold-out uses one fixed split, which is easy to explain. Cross-validation tests every student once, so it depends "
                 "less on luck. We report 5-fold CV as the main result.")

# --- 67 / 68 (hidden)
replace_picture(S[67], pictures(S[67])[0], FIG / "tree_sgpa_large.png")
set_text(find(S[68], start="WEKA Screenshots: Random Forest Runs"), "WEKA Output: Random Forest Runs")
set_text(find(S[68], start="Main CGPA task: Random Forest"), f"Main CGPA task: Random Forest {100 * CV['main_cgpa']['RandomForest']['accuracy']:.1f}% (5-fold CV)")
set_text(find(S[68], start="CGPA questionnaire diagnostic"), f"CGPA questionnaire: Random Forest {100 * CV['questionnaire_cgpa']['RandomForest']['accuracy']:.1f}% (5-fold CV)")
left, right = pictures(S[68])
replace_picture(S[68], left, FIG / "weka_output_main_rf_cv5.png")
replace_picture(S[68], right, FIG / "weka_output_quest_rf_cv5.png")
set_notes(S[68], "Real WEKA 3.8.7 output text. Replace with WEKA Explorer screenshots if you take them.")


# ============================================================================ new slides
def blank_page_number(slide):
    for sh in slide.shapes:
        if texty(sh) and re.fullmatch(r"\d+", sh.text_frame.text.strip()) and sh.top > Pt(480):
            set_text(sh, "")


within = G["gap_distribution_all"]
w01 = int(within["0"]) + int(within["1"])
w23 = int(within["2"]) + int(within["3"])
total = w01 + w23

# N1 — do SGPA and CGPA agree?
s = NEW_AGREE
set_text(find(s, name="TextBox 2"), "DATA QUALITY CHECK")
set_text(find(s, name="TextBox 3"), "Do a Student's SGPA and CGPA Agree?")
replace_picture(s, pictures(s)[0], FIG / "gap_distribution_1102.png")
set_text(find(s, name="TextBox 8"), f"{100 * w01 / total:.0f}%")
set_text(find(s, name="TextBox 9"), f"{w01} within 1 band")
set_text(find(s, name="TextBox 11"), f"{100 * w23 / total:.0f}%")
set_text(find(s, name="TextBox 12"), f"{w23} are 2–3 apart")
set_text(find(s, name="Rounded Rectangle 15"), "CGPA averages many semesters, so it moves slowly")
set_text(find(s, name="SimpleLightSectionLabel"), "DATA CHECK")
blank_page_number(s)
set_notes(s, f"For each of the {total} students with a real SGPA, we compared the recent SGPA group with the CGPA group. "
             f"{w01} students ({100 * w01 / total:.0f}%) are in the same group or one group apart. {w23} students are 2 or 3 "
             "groups apart. Because CGPA is an average of all semesters, a gap that big should be rare.")

# N2 — one semester moves CGPA only a little
s = NEW_EXAMPLE
set_text(find(s, name="TextBox 2"), "DATA QUALITY CHECK")
set_text(find(s, name="TextBox 3"), "One Semester Moves CGPA Only a Little")
set_text(find(s, name="TextBox 8"), "Example: a bad semester")
set_text(find(s, name="TextBox 9"), "• CGPA 3.80 after 5 semesters\n• 6th semester SGPA: 2.80\n• New CGPA: 3.63\n• Group moves only C3 → C2")
set_text(find(s, name="TextBox 12"), "Example: a great semester")
set_text(find(s, name="TextBox 13"), "• CGPA 3.00 after 5 semesters\n• 6th semester SGPA: 4.00\n• New CGPA: 3.17\n• Group stays C0")
set_text(find(s, name="Rounded Rectangle 14"), "Big gaps can happen for a few students. When many answers in one group show them, we check that group.")
set_text(find(s, name="SimpleLightSectionLabel"), "DATA CHECK")
blank_page_number(s)
set_notes(s, "Simple arithmetic: new CGPA = (5 × old CGPA + new SGPA) / 6. A very bad semester (3.80 then 2.80) gives 3.63, "
             "a very good semester (3.00 then 4.00) gives 3.17. So a big SGPA/CGPA gap is possible for a few students, but it "
             "should not be common.")

# N3 — where were the big gaps?
s = NEW_WHERE
set_text(find(s, name="TextBox 2"), "DATA QUALITY CHECK")
set_text(find(s, name="TextBox 3"), "Where Were the Big Gaps?")
set_text(find(s, name="TextBox 6"), "By university (2–3 bands apart)")
set_text(find(s, name="TextBox 9"), "KUET, by semester")
l, r = pictures(s)
replace_picture(s, l, FIG / "far_gap_by_university.png")
replace_picture(s, r, FIG / "far_gap_by_kuet_semester.png")
kuet = next(x for x in G["far_gap_by_university"] if x["university"] == "KUET")
buet = next(x for x in G["far_gap_by_university"] if x["university"] == "BUET")
sem7 = next(x for x in G["far_gap_by_kuet_semester"] if int(x["semester"]) == 7)
set_text(find(s, name="Rounded Rectangle 8"), f"KUET has {kuet['far']} of the {w23} big gaps")
set_text(find(s, name="Rounded Rectangle 11"), f"KUET semester 7: {sem7['far']} of {sem7['n']} ({100 * sem7['far'] / sem7['n']:.0f}%)")
set_text(find(s, name="SimpleLightSectionLabel"), "DATA CHECK")
blank_page_number(s)
set_notes(s, f"BRAC and CUET have no big gaps. BUET has {buet['far']} of {buet['n']}, KUET has {kuet['far']} of {kuet['n']}. "
             f"Inside KUET almost all big gaps are in semester 7 ({sem7['far']} of {sem7['n']}). Other KUET semesters have 0 to 3 each. "
             "BUET was not removed: the numbers are smaller and there is no single time pattern.")

# N4 — when were they submitted?
s = NEW_WHEN
set_text(find(s, name="TextBox 2"), "DATA QUALITY CHECK")
set_text(find(s, name="TextBox 3"), "When Were These Answers Submitted?")
replace_picture(s, pictures(s)[0], FIG / "kuet_sem7_by_date.png")
dates = G["kuet_sem7_by_date"]
wave_n = sum(d["n"] for d in dates if d["date"] in ("2026-08-06", "2026-08-07"))
wave_far = sum(d["far"] for d in dates if d["date"] in ("2026-08-06", "2026-08-07"))
pre = [d for d in dates if d["date"] < "2026-08-06"]
pre_n, pre_far = sum(d["n"] for d in pre), sum(d["far"] for d in pre)
set_text(find(s, name="TextBox 8"), f"{100 * wave_far / wave_n:.0f}%")
set_text(find(s, name="TextBox 9"), f"of the 6–7 Aug answers had a big gap ({wave_far} of {wave_n})")
set_text(find(s, name="SimpleLightSectionLabel"), "DATA CHECK")
blank_page_number(s)
set_notes(s, f"KUET semester-7 answers came in over two weeks. On 6 and 7 August we got {wave_n} answers in two days, and "
             f"{wave_far} of them ({100 * wave_far / wave_n:.0f}%) had a big SGPA/CGPA gap. Before 6 August only {pre_far} of {pre_n} "
             "did. The form is anonymous, so we do not know who filled these answers, only that this group looks different.")

# N5 — the 6-7 Aug answers look random
s = NEW_RANDOM
set_text(find(s, name="TextBox 2"), "DATA QUALITY CHECK")
set_text(find(s, name="TextBox 3"), "The 6–7 August Answers Look Random")
set_text(find(s, name="TextBox 6"), f"6–7 Aug, KUET semester 7 ({G['aug67_group_rows']})")
set_text(find(s, name="TextBox 9"), f"All other students ({G['base_rows'] - G['aug67_group_rows']})")
l, r = pictures(s)
replace_picture(s, l, FIG / "sgpa_x_cgpa_aug67.png")
replace_picture(s, r, FIG / "sgpa_x_cgpa_rest.png")
set_text(find(s, name="Rounded Rectangle 8"), "Answers spread over every cell")
set_text(find(s, name="Rounded Rectangle 11"), "Most answers sit near the diagonal")
set_text(find(s, name="SimpleLightSectionLabel"), "DATA CHECK")
blank_page_number(s)
set_notes(s, "Each table counts students by recent SGPA group (rows) and CGPA group (columns). Blue cells = same group or one "
             f"group apart. For most students the numbers sit near the diagonal (SGPA-CGPA correlation {G['sgpa_cgpa_rho_rest']:.2f}). "
             f"In the 6-7 August group the numbers are spread almost evenly (correlation {G['sgpa_cgpa_rho_aug67']:.2f}), "
             "which is what random picking looks like.")

# N6 — what we kept and what we removed
s = NEW_KEPT
set_text(find(s, name="TextBox 2"), "DATA QUALITY CHECK")
set_text(find(s, name="TextBox 3"), "What We Kept and What We Removed")
set_text(find(s, name="Rounded Rectangle 6"), "1,108 clean rows")
set_text(find(s, name="Rounded Rectangle 8"), "Keep rows with a real SGPA")
set_text(find(s, name="Rounded Rectangle 10"), "Check SGPA vs CGPA")
set_text(find(s, name="Rounded Rectangle 12"), f"{N} final rows")
rows_left = ["6 rows with no real SGPA", f"{G['aug67_group_rows']} answers from 6–7 Aug",
             f"{G['aug67_kept']} with a gap of 0–1 band", f"{G['dropped_rows']} with a gap of 2–3 bands",
             "Result satisfaction question"]
rows_right = ["removed", "checked", "kept", "removed", "not used (subjective)"]
for nm, t in zip(["Rounded Rectangle 15", "Rounded Rectangle 18", "Rounded Rectangle 21", "Rounded Rectangle 24", "Rounded Rectangle 27"], rows_left):
    set_text(find(s, name=nm), t)
for nm, t in zip(["Rounded Rectangle 17", "Rounded Rectangle 20", "Rounded Rectangle 23", "Rounded Rectangle 26", "Rounded Rectangle 29"], rows_right):
    set_text(find(s, name=nm), t)
set_text(find(s, name="SimpleLightSectionLabel"), "DATA CHECK")
blank_page_number(s)
set_notes(s, f"From 1,108 clean rows we kept the {G['base_rows']} with a real recent SGPA. From the {G['aug67_group_rows']} "
             f"answers sent on 6-7 August we kept {G['aug67_kept']} whose SGPA and CGPA are at most one group apart, and removed "
             f"{G['dropped_rows']} with a gap of 2 or 3 groups. That leaves {N} rows, about 83% of the data. We also do not use the "
             "result satisfaction question, because it depends on each student's own goals.")

# N7 — accuracy before and after the data check
s = NEW_BEFORE_AFTER
tb = table_of(s)
set_text(find(s, name="Google Shape;533;p58"), "Accuracy before and after the data check")
set_text(find(s, start="Stratified 10-fold"), "Main CGPA (12 questions + recent SGPA), 5-fold cross-validation, seed 1")
set_cell(tb, 0, 0, "WEKA model")
set_cell(tb, 0, 1, f"All {G['base_rows']:,} students")
set_cell(tb, 0, 2, f"After data check ({N})")
set_cell(tb, 0, 3, "Change")
for r, (label, m) in enumerate([("J48", "J48"), ("Random Forest", "RandomForest"), ("SMO", "SMO"),
                                ("Logistic", "Logistic"), ("Naive Bayes", "NaiveBayes"), ("ZeroR", "ZeroR")], start=1):
    b, a = before[m]["accuracy"], CV["main_cgpa"][m]["accuracy"]
    set_cell(tb, r, 0, label)
    set_cell(tb, r, 1, pct(b))
    set_cell(tb, r, 2, pct(a))
    set_cell(tb, r, 3, f"{100 * (a - b):+.1f} points")
set_text(find(s, start="Same 1,102"), f"Before: {G['base_rows']:,} students with a real SGPA. After: {G['dropped_rows']} answers with a 2–3 band SGPA/CGPA gap removed.")
blank_page_number(s)
set_notes(s, "Same models, same settings, same 5-fold cross-validation. The only change is removing the 182 unreliable answers. "
             f"SMO goes from {pct(before['SMO']['accuracy'])} to {pct(CV['main_cgpa']['SMO']['accuracy'])}. "
             "Logs: docs/weka_runs/final/before1102_main_*_cv5.txt and main_cgpa_*_cv5.txt.")

# ============================================================================ order
anchor = S[9]
for new in [NEW_AGREE, NEW_EXAMPLE, NEW_WHERE, NEW_WHEN, NEW_RANDOM, NEW_KEPT]:
    move_to_after(new, anchor)
    anchor = new
move_to_after(NEW_BEFORE_AFTER, S[48])

prs.save(DECK)
print("saved", DECK.relative_to(ROOT), "| slides:", len(prs.slides))
